import io
import re
from dataclasses import dataclass, field
from typing import List, Optional
from app.core.config import settings


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class DocumentExtractionResult:
    success: bool
    page_count: int
    full_text: str
    pages: List[ExtractedPage] = field(default_factory=list)
    is_scanned: bool = False
    extraction_method: str = "text"
    error_message: Optional[str] = None


def sanitize_extraction_error(error_str: str) -> str:
    """Removes sensitive local filesystem paths from parser exception messages."""
    cleaned = re.sub(r"[A-Za-z]:\\[\w\- \\.]+", "[PATH]", error_str)
    cleaned = re.sub(r"/(?:[a-zA-Z0-9_\-\.]+/)+[a-zA-Z0-9_\-\.]+", "[PATH]", cleaned)
    return cleaned


class DocumentExtractionService:
    """
    Dedicated document extraction service supporting:
    - PDF text extraction via PyMuPDF (fitz)
    - DOCX paragraph and table extraction via python-docx
    - Page-aware tracking
    - Scanned PDF detection
    - Configurable OCR fallback
    - Text normalization
    """

    MIN_TEXT_CHARS_THRESHOLD = 50
    MAX_PDF_PAGES = 50
    MAX_DOCX_PARAGRAPHS = 2000
    MAX_DOCX_TABLE_ROWS = 1000

    def __init__(self, enable_ocr: Optional[bool] = None):
        self.enable_ocr = enable_ocr if enable_ocr is not None else settings.ENABLE_OCR

    def normalize_text(self, text: str) -> str:
        """
        Normalizes document text without destroying meaningful structure like
        lists, line breaks, or section headers.
        """
        if not text:
            return ""
        # Normalize newlines
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        # Replace non-breaking spaces and tabs with normal spaces
        normalized = normalized.replace("\u00a0", " ").replace("\t", "    ")
        # Collapse excessive blank lines (more than 2 newlines) into 2 newlines
        normalized = re.sub(r"\n{3,}", "\n\n", normalized)
        # Strip trailing spaces on each line
        lines = [line.rstrip() for line in normalized.split("\n")]
        return "\n".join(lines).strip()

    def extract_pdf(self, content: bytes) -> DocumentExtractionResult:
        """Extract text from PDF using PyMuPDF with scanned document detection and OCR fallback."""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            return DocumentExtractionResult(
                success=False,
                page_count=0,
                full_text="",
                extraction_method="pymupdf",
                error_message="PyMuPDF library is not installed in runtime environment.",
            )

        try:
            doc = fitz.open(stream=content, filetype="pdf")
            page_count = len(doc)
            if page_count > self.MAX_PDF_PAGES:
                doc.close()
                return DocumentExtractionResult(
                    success=False,
                    page_count=page_count,
                    full_text="",
                    extraction_method="pymupdf",
                    error_message=f"Document exceeds maximum allowed length of {self.MAX_PDF_PAGES} pages (found {page_count} pages).",
                )

            pages: List[ExtractedPage] = []
            total_chars = 0

            for i in range(page_count):
                page = doc.load_page(i)
                text = page.get_text("text") or ""
                cleaned = self.normalize_text(text)
                pages.append(ExtractedPage(page_number=i + 1, text=cleaned))
                total_chars += len(cleaned)

            # Check if document appears to be scanned (very little or no extractable text)
            is_scanned = total_chars < self.MIN_TEXT_CHARS_THRESHOLD

            if is_scanned:
                if self.enable_ocr:
                    # Attempt OCR fallback if enabled
                    ocr_res = self._ocr_fallback_pdf(doc)
                    if ocr_res and ocr_res.success:
                        doc.close()
                        return ocr_res
                # Document is scanned and OCR is either disabled or unavailable
                doc.close()
                return DocumentExtractionResult(
                    success=False,
                    page_count=page_count,
                    full_text="",
                    pages=pages,
                    is_scanned=True,
                    extraction_method="pymupdf_scanned",
                    error_message=(
                        "Scanned document detected with no extractable digital text. "
                        "OCR fallback is currently disabled or unavailable."
                    ),
                )

            full_text = "\n\n--- Page Break ---\n\n".join(
                [f"--- Page {p.page_number} ---\n{p.text}" for p in pages if p.text]
            )

            doc.close()
            return DocumentExtractionResult(
                success=True,
                page_count=page_count,
                full_text=full_text,
                pages=pages,
                is_scanned=False,
                extraction_method="pymupdf",
            )
        except Exception as e:
            return DocumentExtractionResult(
                success=False,
                page_count=0,
                full_text="",
                extraction_method="pymupdf",
                error_message=sanitize_extraction_error(f"Failed to extract PDF document: {str(e)}"),
            )

    def _ocr_fallback_pdf(self, doc) -> Optional[DocumentExtractionResult]:
        """Optional OCR fallback when OCR is enabled in configuration."""
        try:
            import pytesseract
            from PIL import Image

            pages: List[ExtractedPage] = []
            total_text = []

            for i in range(len(doc)):
                page = doc.load_page(i)
                pix = page.get_pixmap(dpi=150)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img)
                cleaned = self.normalize_text(text)
                pages.append(ExtractedPage(page_number=i + 1, text=cleaned))
                if cleaned:
                    total_text.append(f"--- Page {i + 1} ---\n{cleaned}")

            full_text = "\n\n--- Page Break ---\n\n".join(total_text)
            return DocumentExtractionResult(
                success=bool(full_text.strip()),
                page_count=len(doc),
                full_text=full_text,
                pages=pages,
                is_scanned=True,
                extraction_method="pymupdf+ocr_fallback",
            )
        except Exception:
            # OCR failed or dependencies missing
            return None

    def extract_docx(self, content: bytes) -> DocumentExtractionResult:
        """Extract text from DOCX including paragraphs and tables."""
        try:
            from docx import Document
        except ImportError:
            return DocumentExtractionResult(
                success=False,
                page_count=0,
                full_text="",
                extraction_method="python-docx",
                error_message="python-docx library is not installed in runtime environment.",
            )

        try:
            doc = Document(io.BytesIO(content))
            paragraphs = []

            # 1. Paragraphs (bounded to prevent memory/CPU exhaustion)
            for i, p in enumerate(doc.paragraphs):
                if i >= self.MAX_DOCX_PARAGRAPHS:
                    break
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            # 2. Tables (bounded to prevent memory/CPU exhaustion)
            total_rows_processed = 0
            for table in doc.tables:
                if total_rows_processed >= self.MAX_DOCX_TABLE_ROWS:
                    break
                for row in table.rows:
                    total_rows_processed += 1
                    if total_rows_processed > self.MAX_DOCX_TABLE_ROWS:
                        break
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        # De-duplicate adjacent identical cells due to merges
                        deduped = []
                        for rt in row_texts:
                            if not deduped or deduped[-1] != rt:
                                deduped.append(rt)
                        paragraphs.append(" | ".join(deduped))

            raw_text = "\n\n".join(paragraphs)
            cleaned = self.normalize_text(raw_text)

            if len(cleaned) < self.MIN_TEXT_CHARS_THRESHOLD:
                return DocumentExtractionResult(
                    success=False,
                    page_count=1,
                    full_text="",
                    is_scanned=True,
                    extraction_method="python-docx",
                    error_message="Document contains insufficient extractable text.",
                )

            # DOCX does not have fixed physical pages, treat as single logical document
            pages = [ExtractedPage(page_number=1, text=cleaned)]

            return DocumentExtractionResult(
                success=True,
                page_count=1,
                full_text=cleaned,
                pages=pages,
                is_scanned=False,
                extraction_method="python-docx",
            )
        except Exception as e:
            return DocumentExtractionResult(
                success=False,
                page_count=0,
                full_text="",
                extraction_method="python-docx",
                error_message=sanitize_extraction_error(f"Failed to extract DOCX document: {str(e)}"),
            )

    def extract_document(self, content: bytes, mime_type: str) -> DocumentExtractionResult:
        """Unified entry point for document text extraction based on MIME type or content."""
        if mime_type == "application/pdf" or content.startswith(b"%PDF-"):
            return self.extract_pdf(content)
        elif (
            mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or content.startswith(b"PK\x03\x04")
        ):
            return self.extract_docx(content)
        else:
            return DocumentExtractionResult(
                success=False,
                page_count=0,
                full_text="",
                extraction_method="unknown",
                error_message=f"Unsupported document MIME type '{mime_type}'.",
            )
